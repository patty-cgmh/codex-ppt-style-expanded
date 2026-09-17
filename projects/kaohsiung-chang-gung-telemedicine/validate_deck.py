#!/usr/bin/env python3
"""Structural checks for the generated telemedicine PPTX."""

from pathlib import Path
import re

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


ROOT = Path(__file__).resolve().parent
PPTX = ROOT / "output" / "kaohsiung-chang-gung-telemedicine-v1.pptx"
PROMPTS = ROOT / "illustration-prompts.md"


def main():
    assert PPTX.exists(), f"missing {PPTX}"
    prs = Presentation(PPTX)
    assert len(prs.slides) == 15, f"expected 15 slides, got {len(prs.slides)}"
    assert round(prs.slide_width / prs.slide_height, 3) == round(16 / 9, 3)

    all_text = []
    picture_count = 0
    out_of_bounds = []
    for slide_no, slide in enumerate(prs.slides, 1):
        slide_text = []
        for shape in slide.shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                picture_count += 1
            if shape.has_text_frame:
                slide_text.append(shape.text)
            if shape.left < 0 or shape.top < 0 or shape.left + shape.width > prs.slide_width or shape.top + shape.height > prs.slide_height:
                out_of_bounds.append((slide_no, shape.name))
        assert slide_text, f"slide {slide_no} has no editable text"
        assert "〔待補" in "\n".join(slide_text) or slide_no in (1, 7, 8), f"slide {slide_no} lacks explicit data placeholder"
        all_text.extend(slide_text)

    assert not out_of_bounds, f"shapes outside slide: {out_of_bounds}"
    assert picture_count == 0, "v1 should contain native placeholders, not fake raster illustrations"

    deck_ids = set(re.findall(r"IL-\d{2}", "\n".join(all_text)))
    prompt_ids = set(re.findall(r"IL-\d{2}", PROMPTS.read_text(encoding="utf-8")))
    expected = {f"IL-{i:02d}" for i in range(1, 7)}
    assert deck_ids == expected, f"deck illustration IDs: {sorted(deck_ids)}"
    assert prompt_ids == expected, f"prompt illustration IDs: {sorted(prompt_ids)}"
    print(f"PASS: {len(prs.slides)} slides, 16:9, {len(deck_ids)} illustration placeholders, {picture_count} raster pictures")


if __name__ == "__main__":
    main()
