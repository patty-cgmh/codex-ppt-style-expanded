#!/usr/bin/env python3
"""Build the first editable 15-slide Clinical Calm telemedicine deck."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "output"
PPTX = OUT / "kaohsiung-chang-gung-telemedicine-v1.pptx"

SW, SH = 13.333333, 7.5
NAVY, WHITE, MIST = "163A5F", "FFFFFF", "F6F9FC"
ORANGE, TEAL, PALE = "F28C28", "2A8C8B", "DCE8F1"
TEXT, MUTED, LINE = "17324D", "60758A", "C8D7E3"
FONT = "Noto Sans TC"


def rgb(hex_color):
    return RGBColor.from_string(hex_color)


def add_rect(slide, x, y, w, h, fill=WHITE, line=None, radius=True, width=1):
    kind = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb = rgb(fill)
    if line:
        shape.line.color.rgb = rgb(line); shape.line.width = Pt(width)
    else:
        shape.line.fill.background()
    if radius and shape.adjustments:
        shape.adjustments[0] = 0.08
    return shape


def add_text(slide, text, x, y, w, h, size=22, color=TEXT, bold=False,
             align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, margin=0.03):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.clear(); tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    lines = text.split("\n")
    for idx, line in enumerate(lines):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = line; p.alignment = align; p.space_after = Pt(0)
        p.font.name = FONT; p.font.size = Pt(size); p.font.bold = bold; p.font.color.rgb = rgb(color)
    return box


def add_line(slide, x1, y1, x2, y2, color=LINE, width=2, dashed=False):
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    line.line.color.rgb = rgb(color); line.line.width = Pt(width)
    if dashed: line.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    return line


def add_circle(slide, x, y, d, fill=WHITE, line=NAVY, width=2):
    shp = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    shp.fill.solid(); shp.fill.fore_color.rgb = rgb(fill)
    shp.line.color.rgb = rgb(line); shp.line.width = Pt(width)
    return shp


def add_placeholder(slide, code, label, x, y, w, h):
    shp = add_rect(slide, x, y, w, h, fill="F8FBFC", line=TEAL, radius=True, width=1.5)
    shp.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    add_circle(slide, x + 0.18, y + 0.18, 0.34, fill=PALE, line=TEAL, width=1)
    add_text(slide, "+", x + 0.18, y + 0.15, 0.34, 0.35, 17, TEAL, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    add_text(slide, code, x + 0.62, y + 0.20, w - 0.82, 0.32, 12, TEAL, True)
    add_text(slide, label, x + 0.18, y + h/2 - 0.18, w - 0.36, 0.5, 15, MUTED, False, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)


def add_title(slide, number, title, kicker=None):
    add_text(slide, f"{number:02d}", 0.55, 0.45, 0.45, 0.35, 12, ORANGE, True)
    add_text(slide, title, 1.05, 0.32, 11.3, 0.64, 27, NAVY, True, valign=MSO_ANCHOR.MIDDLE)
    if kicker: add_text(slide, kicker, 1.06, 0.96, 11.1, 0.34, 12, MUTED)
    add_line(slide, 0.58, 1.30, 12.75, 1.30, PALE, 1)


def add_footer(slide, n):
    add_text(slide, "高雄長庚遠距醫療服務網絡建置與成果", 0.58, 7.15, 5.8, 0.2, 9, MUTED)
    add_text(slide, f"{n:02d} / 15", 11.8, 7.13, 0.9, 0.22, 9, MUTED, False, PP_ALIGN.RIGHT)


def card(slide, x, y, w, h, title, body="〔待補資料〕", accent=ORANGE):
    add_rect(slide, x, y, w, h, WHITE, PALE)
    add_rect(slide, x, y, 0.08, h, accent, None, False)
    add_text(slide, title, x + 0.25, y + 0.18, w - 0.45, 0.34, 13, MUTED, True)
    add_text(slide, body, x + 0.25, y + 0.62, w - 0.45, h - 0.78, 20, NAVY, True, valign=MSO_ANCHOR.MIDDLE)


def tag(slide, text, x, y, w, fill=PALE, color=NAVY):
    add_rect(slide, x, y, w, 0.38, fill, None)
    add_text(slide, text, x, y + 0.01, w, 0.34, 11, color, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def base_slide(prs, n, title, kicker=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill; bg.solid(); bg.fore_color.rgb = rgb(MIST)
    add_title(slide, n, title, kicker); add_footer(slide, n)
    return slide


def build():
    OUT.mkdir(parents=True, exist_ok=True)
    prs = Presentation(); prs.slide_width = Inches(SW); prs.slide_height = Inches(SH)

    # 1 Cover
    s = prs.slides.add_slide(prs.slide_layouts[6]); s.background.fill.solid(); s.background.fill.fore_color.rgb = rgb(MIST)
    add_rect(s, 0, 0, 0.17, SH, ORANGE, None, False)
    tag(s, "醫療行政成果報告", 0.78, 0.64, 1.75, WHITE, NAVY)
    add_text(s, "高雄長庚\n遠距醫療服務網絡\n建置與成果", 0.78, 1.35, 6.35, 2.55, 30, NAVY, True)
    add_text(s, "以可信賴的跨域協作，延伸可近的醫療服務", 0.82, 4.20, 5.85, 0.5, 17, TEAL, True)
    add_text(s, "報告單位｜〔待補資料〕   報告人｜〔待補資料〕   日期｜〔待補資料〕", 0.82, 6.65, 7.4, 0.3, 11, MUTED)
    add_placeholder(s, "IL-01", "友善線性醫療插圖｜遠距照護連線", 8.25, 1.25, 4.25, 4.85)
    notes(s, "封面。所有單位、報告人與日期均待院方確認；右側 IL-01 將替換為正式友善線性醫療插圖。")

    # 2 Executive summary
    s = base_slide(prs, 2, "本期成果以網絡、服務、協作與品質四個面向呈現", "執行摘要｜數字取得前保留明確資料缺口")
    for i, (t, c) in enumerate([("網絡覆蓋", NAVY), ("服務量能", ORANGE), ("跨域協作", TEAL), ("品質成果", NAVY)]):
        x = 0.72 + i*3.08; card(s, x, 1.72, 2.65, 2.15, t, "〔待補資料〕", c)
    add_rect(s, 0.72, 4.30, 11.89, 1.55, WHITE, PALE)
    add_text(s, "本頁核心結論", 1.02, 4.62, 1.55, 0.32, 12, ORANGE, True)
    add_text(s, "待核心 KPI、統計期間與比較基準確認後，形成可供主管快速決策的成果全貌。", 2.55, 4.48, 9.45, 0.65, 20, NAVY, True, valign=MSO_ANCHOR.MIDDLE)
    add_text(s, "資料規則｜不以示意數字取代正式成果", 1.02, 5.40, 6.8, 0.28, 11, MUTED)
    notes(s, "填入經核實的四項核心成果；每項必須附統計期間與口徑。")

    # 3 Background
    s = base_slide(prs, 3, "網絡建置回應醫療可近性與跨院協作的實際需求", "背景與需求｜先說明問題，再連結建置策略")
    card(s, 0.72, 1.72, 3.25, 1.35, "服務可近性", "需求證據〔待補資料〕", NAVY)
    card(s, 0.72, 3.30, 3.25, 1.35, "跨院協作", "需求證據〔待補資料〕", TEAL)
    card(s, 0.72, 4.88, 3.25, 1.35, "照護連續性", "需求證據〔待補資料〕", ORANGE)
    add_line(s, 4.25, 3.98, 5.20, 3.98, ORANGE, 3)
    add_text(s, "→", 4.48, 3.72, 0.5, 0.45, 22, ORANGE, True, PP_ALIGN.CENTER)
    add_placeholder(s, "IL-02", "友善線性醫療插圖｜跨距離照護", 5.35, 1.72, 3.25, 4.52)
    add_rect(s, 8.88, 1.72, 3.72, 4.52, WHITE, PALE)
    add_text(s, "建置回應", 9.18, 2.05, 3.0, 0.35, 13, ORANGE, True)
    for i, txt in enumerate(["服務模式〔待補資料〕", "協作機制〔待補資料〕", "品質管理〔待補資料〕"]):
        tag(s, txt, 9.18, 2.72+i*0.88, 3.05, PALE if i != 1 else "E4F0EF", NAVY)
    notes(s, "需補上需求來源，例如正式計畫、院內分析或政策依據；插圖不可暗示尚未證實的成效。")

    # 4 Goals
    s = base_slide(prs, 4, "專案以服務、協作與品質三類目標同步推進", "建置目標｜目標須對應可衡量指標")
    goals = [("01", "服務", "服務對象與可近性\n〔待補資料〕", NAVY), ("02", "協作", "跨域分工與連結\n〔待補資料〕", TEAL), ("03", "品質", "安全、效率與體驗\n〔待補資料〕", ORANGE)]
    for i, (num, title, body, c) in enumerate(goals):
        x = 0.78 + i*4.10; add_rect(s, x, 1.72, 3.62, 4.45, WHITE, PALE)
        add_circle(s, x+0.30, 2.02, 0.62, c, c); add_text(s, num, x+0.30, 2.08, 0.62, 0.38, 13, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        add_text(s, title, x+0.30, 2.93, 2.9, 0.48, 24, NAVY, True)
        add_text(s, body, x+0.30, 3.66, 2.9, 1.02, 17, TEXT)
        add_line(s, x+0.30, 5.12, x+3.10, 5.12, PALE, 1)
        add_text(s, "衡量指標｜〔待補資料〕", x+0.30, 5.35, 2.9, 0.38, 12, MUTED)
    notes(s, "以核定計畫目標替換待補資料，避免自行擴寫專案承諾。")

    # 5 Timeline
    s = base_slide(prs, 5, "制度、系統與服務布建依階段逐步完成", "建置歷程｜所有年月與里程碑待核實")
    add_line(s, 1.05, 3.75, 12.20, 3.75, PALE, 5)
    stages = [(1.15, "階段一", "規劃與治理"), (4.05, "階段二", "系統與流程"), (6.95, "階段三", "試行與擴展"), (9.85, "階段四", "成果與精進")]
    for i, (x, lab, title) in enumerate(stages):
        c = ORANGE if i == 3 else TEAL if i == 2 else NAVY
        add_circle(s, x, 3.40, 0.70, c, WHITE, 3)
        tag(s, lab, x-0.20, 2.15, 1.10, c, WHITE)
        add_text(s, title, x-0.55, 4.35, 1.85, 0.45, 15, NAVY, True, PP_ALIGN.CENTER)
        add_text(s, "年月〔待補資料〕\n里程碑〔待補資料〕", x-0.65, 4.92, 2.05, 0.76, 12, MUTED, False, PP_ALIGN.CENTER)
    notes(s, "四階段僅為版面骨架；正式名稱、時間與成果由院方資料替換。")

    # 6 Network
    s = base_slide(prs, 6, "高雄長庚串聯合作院所，逐步形成遠距醫療服務網絡", "服務網絡｜院所名稱、位置與連線關係待核實")
    add_rect(s, 0.72, 1.62, 8.14, 4.92, WHITE, PALE)
    add_text(s, "服務網絡地圖｜資料待補", 1.02, 1.92, 3.2, 0.34, 13, MUTED, True)
    cx, cy = 4.72, 4.02
    for nx, ny in [(1.55,2.85),(2.05,5.35),(6.92,2.62),(7.35,4.88),(5.95,5.55)]:
        add_line(s, cx, cy, nx+0.25, ny+0.25, TEAL, 2, True); add_circle(s, nx, ny, 0.50, WHITE, ORANGE, 2)
    add_circle(s, cx-0.42, cy-0.42, 0.84, NAVY, NAVY, 2)
    add_text(s, "高雄\n長庚", cx-0.40, cy-0.31, 0.80, 0.58, 12, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    add_text(s, "合作院所節點與地理位置〔待補資料〕", 1.08, 6.02, 6.9, 0.30, 11, MUTED)
    card(s, 9.18, 1.62, 3.42, 1.34, "合作院所數", "〔待補資料〕", ORANGE)
    card(s, 9.18, 3.18, 3.42, 1.34, "涵蓋區域", "〔待補資料〕", NAVY)
    card(s, 9.18, 4.74, 3.42, 1.34, "服務科別", "〔待補資料〕", TEAL)
    notes(s, "中央與周邊節點為資料 placeholder，不代表真實院所數或位置；正式版須依核實資料重建。")

    # 7 Collaboration
    s = base_slide(prs, 7, "跨域角色透過明確分工共同完成遠距照護", "合作模式｜角色與治理機制待院方確認")
    add_placeholder(s, "IL-03", "友善線性醫療插圖｜跨域團隊", 0.72, 1.70, 3.35, 4.78)
    roles = [("高雄長庚", "專業支援與治理", NAVY), ("合作院所", "在地照護與執行", ORANGE), ("協調團隊", "排程、轉介與追蹤", TEAL), ("行政／資訊", "制度與系統支援", NAVY)]
    for i,(r,b,c) in enumerate(roles):
        x=4.42+(i%2)*4.05; y=1.70+(i//2)*2.38; card(s,x,y,3.62,1.92,r,b+"\n〔待確認〕",c)
    add_text(s, "治理原則｜權責清楚・資訊安全・閉環追蹤", 4.50, 6.10, 7.3, 0.34, 14, NAVY, True, PP_ALIGN.CENTER)
    notes(s, "角色名稱為建議架構；正式職責與治理關係須依實際作業文件調整。")

    # 8 Process
    s = base_slide(prs, 8, "個案由啟動、會診到追蹤形成閉環服務流程", "服務流程｜起訖點、責任與例外路徑待確認")
    lanes=[("合作院所",1.70,"F1F6FA"),("高雄長庚",3.15,"FFFFFF"),("共同追蹤",4.60,"F1F6FA")]
    for name,y,fc in lanes:
        add_rect(s,0.72,y,11.9,1.20,fc,PALE); tag(s,name,0.92,y+0.38,1.35,NAVY,WHITE)
    steps=[(2.65,2.00,"需求啟動",ORANGE),(4.55,3.45,"資料確認",NAVY),(6.45,3.45,"遠距會診",TEAL),(8.35,4.90,"建議執行",NAVY),(10.25,4.90,"追蹤結案",ORANGE)]
    for i,(x,y,t,c) in enumerate(steps):
        add_rect(s,x,y,1.42,0.62,WHITE,c); add_text(s,t,x,y+0.12,1.42,0.32,12,c,True,PP_ALIGN.CENTER,MSO_ANCHOR.MIDDLE)
        if i<len(steps)-1: add_line(s,x+1.42,y+0.31,steps[i+1][0],steps[i+1][1]+0.31,TEAL,2)
    add_text(s,"判斷條件與例外路徑〔待補資料〕",4.62,6.20,4.8,0.30,11,MUTED,False,PP_ALIGN.CENTER)
    notes(s, "此頁用原生 PowerPoint 形狀建立，可直接修改節點與連線；現有步驟是結構 placeholder。")

    # 9 Services
    s = base_slide(prs, 9, "服務場景將依對象、科別與工具清楚分層", "服務內容｜不預設尚未提供的科別或服務項目")
    add_rect(s,0.72,1.68,8.22,4.88,WHITE,PALE)
    headers=[("服務對象",1.02,2.15),("服務科別",3.56,2.15),("服務項目",6.10,2.15)]
    for t,x,y in headers: tag(s,t,x,y,2.12,NAVY,WHITE)
    for row in range(3):
        yy=3.05+row*0.90
        for col,x in enumerate([1.02,3.56,6.10]): tag(s,"〔待補資料〕",x,yy,2.12,"F1F6FA",MUTED)
    add_placeholder(s,"IL-04","友善線性醫療插圖｜設備與照護",9.22,1.68,3.38,4.88)
    notes(s, "矩陣不填入推測的科別；待取得正式服務清單後再調整行列。")

    # 10 KPI
    s = base_slide(prs, 10, "成果應以一致口徑的 KPI 集中檢視", "成果總覽｜每一指標須附定義、單位與期間")
    kpis=[("服務量",ORANGE),("合作院所",NAVY),("服務科別",TEAL),("完成率",NAVY),("回應時效",TEAL),("滿意度",ORANGE)]
    for i,(t,c) in enumerate(kpis):
        x=0.72+(i%3)*4.00; y=1.70+(i//3)*2.27
        card(s,x,y,3.55,1.86,t,"〔待補資料〕",c); add_text(s,"統計期間／定義〔待補〕",x+0.25,y+1.46,2.95,0.22,10,MUTED)
    add_rect(s,0.72,6.33,11.55,0.42,PALE,None); add_text(s,"解讀原則｜不比較不同口徑、不省略母數、不以百分比取代實際量",0.92,6.39,11.1,0.25,11,NAVY,True,PP_ALIGN.CENTER)
    notes(s, "六張卡為可調整槽位；依實際 KPI 數量增減，不用示意數字。")

    # 11 Trend
    s = base_slide(prs, 11, "服務使用趨勢需搭配事件與統計口徑解讀", "服務趨勢｜原始時序資料待補")
    add_rect(s,0.72,1.70,8.42,4.88,WHITE,PALE)
    add_text(s,"服務趨勢圖",1.02,1.98,2.4,0.34,13,MUTED,True)
    for i in range(5): add_line(s,1.12,2.65+i*0.68,8.72,2.65+i*0.68,"E4ECF2",1)
    add_line(s,1.12,5.70,8.72,5.70,LINE,1.5); add_line(s,1.12,2.55,1.12,5.70,LINE,1.5)
    add_text(s,"〔待補時序原始資料〕",2.62,3.78,4.6,0.52,20,PALE,True,PP_ALIGN.CENTER)
    add_text(s,"時間",7.95,5.85,0.75,0.25,10,MUTED,False,PP_ALIGN.RIGHT)
    add_rect(s,9.45,1.70,3.15,4.88,WHITE,PALE)
    add_text(s,"解讀卡",9.75,2.02,2.4,0.32,13,ORANGE,True)
    for i,t in enumerate(["趨勢變化〔待補〕","關鍵事件〔待補〕","比較基準〔待補〕"]):
        tag(s,t,9.75,2.72+i*1.00,2.38,PALE if i!=1 else "FFF2E5",NAVY)
    notes(s, "正式資料匯入後再建立原生圖表；目前只保留座標與解讀區，不繪製虛構曲線。")

    # 12 Heatmap
    s = base_slide(prs, 12, "服務分布可從地區、院所、科別或時段辨識差異", "服務分布｜熱力維度、分母與原始值待補")
    add_rect(s,0.72,1.70,8.42,4.90,WHITE,PALE)
    add_text(s,"服務分布熱力表",1.02,1.98,2.8,0.34,13,MUTED,True)
    for r in range(4):
        add_text(s,f"維度 {r+1}",1.02,2.72+r*0.72,1.18,0.30,11,MUTED)
        for c in range(6):
            fill=["EFF5F8","E4EEF3","D8E8EE"][((r+c)%3)]
            add_rect(s,2.28+c*0.96,2.63+r*0.72,0.78,0.50,fill,WHITE,False)
    add_text(s,"所有色階目前僅為版面 placeholder，不代表服務量",2.30,5.78,5.75,0.28,10,MUTED)
    add_rect(s,9.45,1.70,3.15,4.90,WHITE,PALE)
    add_text(s,"關鍵洞察",9.75,2.02,2.4,0.32,13,ORANGE,True)
    add_text(s,"分布差異\n〔待補資料〕",9.75,2.80,2.38,0.92,19,NAVY,True)
    add_line(s,9.75,4.00,12.05,4.00,PALE,1)
    add_text(s,"分母與口徑\n〔待補資料〕",9.75,4.28,2.38,0.80,14,MUTED,True)
    notes(s, "熱力格目前不代表數值；正式版須依實際維度與原始資料重建單向藍色階。")

    # 13 Highlights
    s = base_slide(prs, 13, "特色亮點必須同時呈現做法與可核實價值", "特色亮點｜不以插圖或口號取代證據")
    add_placeholder(s,"IL-05","友善線性醫療插圖｜人本連結",0.72,1.70,3.42,4.85)
    highlights=[("亮點做法 01","做法〔待補資料〕\n佐證〔待補資料〕",NAVY),("亮點做法 02","做法〔待補資料〕\n佐證〔待補資料〕",TEAL),("公開案例","情境〔待補資料〕\n成效〔待補資料〕",ORANGE)]
    for i,(t,b,c) in enumerate(highlights): card(s,4.48,1.70+i*1.63,8.12,1.38,t,b,c)
    notes(s, "案例須先確認可公開範圍，避免個資、病歷與不可驗證的效果敘述。")

    # 14 Challenges
    s = base_slide(prs, 14, "已知挑戰需要對應具體改善行動與責任", "挑戰與精進｜建立可追蹤的改善閉環")
    add_text(s,"挑戰",0.95,1.72,2.75,0.38,14,NAVY,True)
    add_text(s,"改善行動",4.38,1.72,3.25,0.38,14,NAVY,True)
    add_text(s,"責任與時程",9.05,1.72,2.75,0.38,14,NAVY,True)
    for i in range(3):
        y=2.32+i*1.30
        card(s,0.72,y,3.12,0.98,f"挑戰 {i+1}","〔待補資料〕",ORANGE)
        add_line(s,3.94,y+0.49,4.25,y+0.49,TEAL,2)
        card(s,4.32,y,4.02,0.98,"對應行動","〔待補資料〕",TEAL)
        add_line(s,8.44,y+0.49,8.75,y+0.49,TEAL,2)
        card(s,8.82,y,3.78,0.98,"Owner／期限","〔待補資料〕",NAVY)
    add_rect(s,0.72,6.35,11.88,0.42,PALE,None); add_text(s,"排序原則｜影響程度 × 執行急迫性 × 可行性〔待確認〕",0.95,6.41,11.3,0.25,11,NAVY,True,PP_ALIGN.CENTER)
    notes(s, "只列出院方確認的挑戰；避免用推測問題替代實際復盤。")

    # 15 Roadmap
    s = base_slide(prs, 15, "下一階段以可執行路線圖擴大服務與影響", "下一階段｜目標、時程、資源與 KPI 待補")
    add_rect(s,0.72,1.70,8.18,4.82,WHITE,PALE)
    add_line(s,1.30,4.00,8.25,4.00,PALE,5)
    roadmap=[(1.32,"近期","穩定與補強",NAVY),(3.75,"中期","擴展與整合",TEAL),(6.18,"後期","深化與評估",ORANGE)]
    for x,stage,title,c in roadmap:
        add_circle(s,x,3.62,0.76,c,WHITE,3); tag(s,stage,x-0.05,2.28,0.92,c,WHITE)
        add_text(s,title,x-0.45,4.66,1.68,0.36,14,NAVY,True,PP_ALIGN.CENTER)
        add_text(s,"行動／時程\n〔待補資料〕",x-0.55,5.15,1.90,0.66,11,MUTED,False,PP_ALIGN.CENTER)
    add_placeholder(s,"IL-06","友善線性醫療插圖｜網絡延伸",9.22,1.70,3.38,3.22)
    add_rect(s,9.22,5.18,3.38,1.34,WHITE,PALE)
    add_text(s,"成功指標",9.52,5.42,2.75,0.30,12,ORANGE,True)
    add_text(s,"〔待補資料〕",9.52,5.84,2.75,0.34,18,NAVY,True)
    notes(s, "結尾以三階段路線圖收束；所有目標、資源與 KPI 均待院方核定資料。")

    prs.save(PPTX)
    print(PPTX)


if __name__ == "__main__":
    build()

