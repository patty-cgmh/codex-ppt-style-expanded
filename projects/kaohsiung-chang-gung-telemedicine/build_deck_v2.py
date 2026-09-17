#!/usr/bin/env python3
"""Build the visually redesigned Clinical Calm V2 deck with native vector art."""

from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from build_deck import (
    FONT, LINE, MIST, MUTED, NAVY, ORANGE, PALE, SH, SW, TEAL, TEXT, WHITE,
    add_circle, add_line, add_rect, add_text, notes, rgb,
)


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "output"
PPTX = OUT / "kaohsiung-chang-gung-telemedicine-v2.pptx"
SOFT_TEAL = "E5F2F0"
SOFT_ORANGE = "FFF1E4"
SOFT_BLUE = "EAF2F8"


def oval(slide, x, y, w, h, fill=PALE, line=None, width=1):
    s = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb = rgb(fill)
    if line:
        s.line.color.rgb = rgb(line); s.line.width = Pt(width)
    else: s.line.fill.background()
    return s


def page(prs, n, title, eyebrow):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid(); s.background.fill.fore_color.rgb = rgb(MIST)
    add_text(s, eyebrow.upper(), .65, .38, 3.2, .25, 9, TEAL, True)
    add_text(s, title, .65, .72, 11.8, .54, 25, NAVY, True, valign=MSO_ANCHOR.MIDDLE)
    add_text(s, f"{n:02d}", 12.15, .43, .52, .25, 10, ORANGE, True, PP_ALIGN.RIGHT)
    add_text(s, "高雄長庚｜遠距醫療服務網絡", .65, 7.18, 4.0, .18, 8, MUTED)
    return s


def pill(slide, text, x, y, w, fill=SOFT_TEAL, color=TEAL):
    add_rect(slide, x, y, w, .34, fill, None)
    add_text(slide, text, x, y+.015, w, .28, 10, color, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)


def native_person(slide, x, y, scale=1.0, coat=False, accent=TEAL):
    # Friendly thin-line person; all parts remain native PowerPoint shapes.
    add_circle(slide, x+.34*scale, y, .34*scale, MIST, NAVY, 1.2)
    add_line(slide, x+.51*scale, y+.34*scale, x+.51*scale, y+1.05*scale, NAVY, 1.4)
    add_line(slide, x+.51*scale, y+.52*scale, x+.08*scale, y+.82*scale, NAVY, 1.4)
    add_line(slide, x+.51*scale, y+.52*scale, x+.94*scale, y+.82*scale, NAVY, 1.4)
    add_line(slide, x+.51*scale, y+1.05*scale, x+.18*scale, y+1.48*scale, NAVY, 1.4)
    add_line(slide, x+.51*scale, y+1.05*scale, x+.84*scale, y+1.48*scale, NAVY, 1.4)
    if coat:
        add_line(slide, x+.25*scale, y+.48*scale, x+.18*scale, y+1.15*scale, accent, 2)
        add_line(slide, x+.77*scale, y+.48*scale, x+.84*scale, y+1.15*scale, accent, 2)
        add_circle(slide, x+.46*scale, y+.67*scale, .10*scale, WHITE, ORANGE, 1)


def native_monitor(slide, x, y, w=1.75, h=1.05):
    add_rect(slide, x, y, w, h, WHITE, TEAL, radius=True, width=1.2)
    add_circle(slide, x+w*.41, y+.18, .32, MIST, NAVY, 1)
    add_line(slide, x+w*.50, y+.50, x+w*.50, y+.82, NAVY, 1.2)
    add_line(slide, x+w*.50, y+.60, x+w*.31, y+.75, NAVY, 1.2)
    add_line(slide, x+w*.50, y+.60, x+w*.69, y+.75, NAVY, 1.2)
    add_line(slide, x+w*.50, y+h, x+w*.50, y+h+.25, TEAL, 1.2)
    add_line(slide, x+w*.30, y+h+.25, x+w*.70, y+h+.25, TEAL, 1.2)


def native_building(slide, x, y, w=1.6, h=1.4, remote=False):
    add_rect(slide, x, y+.2, w, h-.2, WHITE, NAVY, radius=True, width=1.2)
    add_rect(slide, x+w*.43, y+.48, w*.14, .42, TEAL, None, False)
    add_rect(slide, x+w*.31, y+.62, w*.38, .12, TEAL, None, False)
    add_rect(slide, x+w*.43, y+h-.25, w*.20, .25, PALE, None, False)
    if remote:
        oval(slide, x-.18, y+h-.05, w+1.0, .30, SOFT_TEAL)


def native_network(slide, cx, cy, radius=1.7, nodes=6):
    import math
    for i in range(nodes):
        a=2*math.pi*i/nodes; x=cx+math.cos(a)*radius; y=cy+math.sin(a)*radius*.62
        add_line(slide,cx,cy,x,y,TEAL,1.2,True); add_circle(slide,x-.14,y-.14,.28,WHITE,ORANGE if i%3==0 else TEAL,1.3)
    add_circle(slide,cx-.34,cy-.34,.68,NAVY,NAVY,1)
    add_text(slide,"醫療\n中心",cx-.31,cy-.25,.62,.48,9,WHITE,True,PP_ALIGN.CENTER,MSO_ANCHOR.MIDDLE)


def metric(slide, x, y, label, size=30, accent=ORANGE, sub="統計期間〔待補〕"):
    add_text(slide,"〔待補資料〕",x,y,2.55,.55,size,NAVY,True)
    add_line(slide,x,y+.68,x+.42,y+.68,accent,3)
    add_text(slide,label,x+.55,y+.55,1.95,.30,11,TEXT,True)
    add_text(slide,sub,x+.55,y+.90,1.95,.24,9,MUTED)


def build():
    OUT.mkdir(parents=True, exist_ok=True)
    prs=Presentation(); prs.slide_width=Inches(SW); prs.slide_height=Inches(SH)

    # 01 Hero
    s=prs.slides.add_slide(prs.slide_layouts[6]); s.background.fill.solid(); s.background.fill.fore_color.rgb=rgb(MIST)
    oval(s,8.10,-1.05,6.3,6.3,SOFT_TEAL); oval(s,9.30,4.95,4.7,3.1,SOFT_BLUE)
    add_text(s,"醫療行政成果報告",.75,.62,2.1,.28,10,TEAL,True)
    add_text(s,"高雄長庚\n遠距醫療服務網絡\n建置與成果",.75,1.38,6.65,2.28,31,NAVY,True)
    add_line(s,.78,4.05,1.48,4.05,ORANGE,4)
    add_text(s,"以可信賴的跨域協作，延伸可近的醫療服務",1.67,3.83,5.2,.40,15,TEAL,True)
    native_person(s,8.35,2.12,1.28,True); native_monitor(s,9.95,2.05,2.05,1.20); native_person(s,11.70,3.95,.80,False)
    add_line(s,9.57,2.85,9.92,2.70,ORANGE,1.5,True); add_line(s,11.90,3.40,12.02,3.92,ORANGE,1.5,True)
    add_text(s,"報告單位｜〔待補資料〕   報告人｜〔待補資料〕   日期｜〔待補資料〕",.78,6.78,7.2,.28,10,MUTED)
    notes(s,"V2 Hero。右側 IL-01 已以 PowerPoint 原生 thin-line 醫師、視訊螢幕與遠端病人插畫落實。")

    # 02 Hero KPI, intentionally not six equal cards
    s=page(prs,2,"本期成果以網絡、服務、協作與品質四個面向呈現","Executive summary")
    oval(s,-1.2,4.85,5.0,3.2,SOFT_BLUE)
    add_text(s,"一眼看見\n網絡成果",.72,1.72,3.4,1.15,27,NAVY,True)
    add_text(s,"數據確認後，將以一個主指標帶領三個支持指標，建立主管決策所需的成果全貌。",.74,3.18,3.15,1.08,14,TEXT)
    add_rect(s,4.38,1.65,4.05,4.55,WHITE,None)
    pill(s,"HERO KPI",4.75,2.00,1.05,SOFT_ORANGE,ORANGE); metric(s,4.75,2.78,"服務量能",36,ORANGE)
    add_text(s,"核心結果的意義與比較基準〔待補資料〕",4.76,4.45,3.10,.74,14,TEXT,True)
    for i,(lab,c) in enumerate([("網絡覆蓋",TEAL),("跨域協作",NAVY),("品質成果",ORANGE)]): metric(s,8.88,1.78+i*1.48,lab,19,c)
    notes(s,"以 Hero KPI + secondary metrics 建立視覺層級，所有數值維持待補。")

    # 03 editorial + IL-02
    s=page(prs,3,"網絡建置回應醫療可近性與跨院協作的實際需求","Background")
    add_text(s,"距離不應成為\n專科醫療的門檻",.72,1.65,4.35,1.15,27,NAVY,True)
    add_text(s,"需求證據〔待補資料〕",.74,3.03,3.0,.34,12,ORANGE,True)
    add_text(s,"服務可近性、跨院協作與照護連續性，構成本計畫的三個需求視角。",.74,3.58,3.72,.92,15,TEXT)
    for i,t in enumerate(["服務可近性","跨院協作","照護連續性"]): pill(s,t,.74,5.00+i*.48,1.85,SOFT_TEAL if i!=1 else SOFT_BLUE,NAVY)
    oval(s,5.20,1.42,7.65,5.25,SOFT_TEAL)
    native_building(s,5.90,3.30,1.72,1.55); native_building(s,10.56,3.42,1.48,1.33,True)
    for i in range(3): add_line(s,7.70,3.75+i*.22,10.35,3.75+i*.22,ORANGE if i==1 else TEAL,1.2,True)
    add_text(s,"醫院",6.35,5.10,.80,.24,10,NAVY,True,PP_ALIGN.CENTER); add_text(s,"偏鄉／離島",10.58,5.10,1.48,.24,10,NAVY,True,PP_ALIGN.CENTER)
    notes(s,"IL-02 使用原生醫院、偏鄉據點與柔和連線；不暗示未證實成果。")

    # 04 pillars
    s=page(prs,4,"專案以服務、協作與品質三類目標同步推進","Objectives")
    add_text(s,"三個策略支柱，共同支撐可持續的遠距照護。",.72,1.55,6.5,.42,16,TEXT)
    for i,(n,t,desc,c) in enumerate([("01","服務","服務對象與可近性\n〔待補資料〕",NAVY),("02","協作","跨域分工與連結\n〔待補資料〕",TEAL),("03","品質","安全、效率與體驗\n〔待補資料〕",ORANGE)]):
        x=.82+i*4.15; oval(s,x,2.15,3.50,3.50,[SOFT_BLUE,SOFT_TEAL,SOFT_ORANGE][i]); add_text(s,n,x+.35,2.55,.55,.30,11,c,True); add_text(s,t,x+.35,3.25,2.6,.45,24,NAVY,True); add_text(s,desc,x+.35,4.12,2.55,.75,14,TEXT); add_text(s,"衡量指標〔待補〕",x+.35,5.18,2.55,.28,10,MUTED)
    notes(s,"以三個無框有機圓形取代等重卡片，內容架構不變。")

    # 05 timeline
    s=page(prs,5,"制度、系統與服務布建依階段逐步完成","Journey")
    add_line(s,.92,3.58,12.20,3.58,TEAL,2)
    stages=[("規劃與治理",NAVY),("系統與流程",TEAL),("試行與擴展",TEAL),("成果與精進",ORANGE)]
    for i,(t,c) in enumerate(stages):
        x=1.05+i*3.05; add_circle(s,x,3.22,.72,MIST,c,2); add_text(s,f"0{i+1}",x,3.42,.72,.20,10,c,True,PP_ALIGN.CENTER); add_text(s,t,x-.35,2.12,1.48,.44,14,NAVY,True,PP_ALIGN.CENTER); add_text(s,"年月〔待補〕\n里程碑〔待補資料〕",x-.62,4.42,2.05,.70,11,MUTED,False,PP_ALIGN.CENTER)
    oval(s,9.88,5.55,3.7,1.7,SOFT_ORANGE)
    notes(s,"水平旅程頁，以單一細線和節點建立節奏。")

    # 06 map-led
    s=page(prs,6,"高雄長庚串聯合作院所，逐步形成遠距醫療服務網絡","Service network")
    oval(s,.42,1.35,8.55,5.52,SOFT_BLUE); oval(s,1.12,2.08,6.55,4.05,WHITE)
    native_network(s,4.45,4.02,2.45,8)
    add_text(s,"示意網絡｜正式地理位置與院所名稱待補",1.25,6.10,6.2,.28,10,MUTED)
    add_text(s,"從單點會診\n走向協作網絡",9.35,1.82,2.85,1.10,24,NAVY,True)
    metric(s,9.38,3.35,"合作院所數",21,ORANGE); metric(s,9.38,4.68,"涵蓋區域",21,NAVY); metric(s,9.38,5.98,"服務科別",18,TEAL,"口徑〔待補〕")
    notes(s,"Map-led hero，以地理感有機底形和非均質節點呈現；節點不代表真實院所數。")

    # 07 human-centered team
    s=page(prs,7,"跨域角色透過明確分工共同完成遠距照護","Collaboration")
    oval(s,-.65,1.48,6.75,5.65,SOFT_TEAL)
    for i,(lab,coat,c) in enumerate([("醫師",True,NAVY),("護理",True,TEAL),("合作院所",False,ORANGE),("行政／資訊",False,NAVY)]):
        x=.58+i*1.28; native_person(s,x,2.48,.84,coat,c); add_text(s,lab,x-.08,4.02,1.18,.30,10,NAVY,True,PP_ALIGN.CENTER)
    add_line(s,1.10,4.70,4.78,4.70,TEAL,1.5); add_circle(s,2.74,4.53,.34,WHITE,ORANGE,1.5)
    add_text(s,"協作不是角色並列，\n而是責任相互銜接。",6.72,1.78,5.10,1.08,25,NAVY,True)
    for i,(a,b) in enumerate([("專業支援與治理","高雄長庚"),("在地照護與執行","合作院所"),("排程、轉介與追蹤","協調團隊"),("制度與系統支援","行政／資訊")]): add_text(s,a,6.75,3.35+i*.72,2.75,.30,13,TEXT,True); add_text(s,b+"〔待確認〕",9.60,3.35+i*.72,2.45,.30,11,MUTED)
    notes(s,"IL-03 以四名原生線性人物呈現跨域協作，避免 icon collage。")

    # 08 process
    s=page(prs,8,"個案由啟動、會診到追蹤形成閉環服務流程","Workflow")
    lanes=[("合作院所",2.05,SOFT_BLUE),("高雄長庚",3.47,WHITE),("共同追蹤",4.89,SOFT_TEAL)]
    for name,y,c in lanes: add_rect(s,.72,y,11.88,1.10,c,None); add_text(s,name,.95,y+.34,1.25,.26,10,NAVY,True)
    steps=[(2.42,2.28,"需求啟動",ORANGE),(4.28,3.70,"資料確認",NAVY),(6.18,3.70,"遠距會診",TEAL),(8.08,5.12,"建議執行",NAVY),(10.02,5.12,"追蹤結案",ORANGE)]
    for i,(x,y,t,c) in enumerate(steps):
        add_circle(s,x,y,.56,WHITE,c,1.5); add_text(s,str(i+1),x,y+.16,.56,.18,9,c,True,PP_ALIGN.CENTER); add_text(s,t,x-.35,y+.78,1.28,.28,11,NAVY,True,PP_ALIGN.CENTER)
        if i<len(steps)-1: add_line(s,x+.56,y+.28,steps[i+1][0],steps[i+1][1]+.28,TEAL,1.5)
    notes(s,"原生泳道與節點；判斷條件、例外路徑待補。")

    # 09 services + equipment illustration
    s=page(prs,9,"服務場景將依對象、科別與工具清楚分層","Services")
    add_text(s,"服務矩陣",.72,1.62,2.0,.34,15,NAVY,True)
    for i,h in enumerate(["服務對象","服務科別","服務項目"]): add_text(s,h,.85+i*2.35,2.25,1.8,.28,11,TEAL,True)
    for r in range(3):
        add_line(s,.78,2.92+r*.92,7.55,2.92+r*.92,PALE,1)
        for c in range(3): add_text(s,"〔待補資料〕",.85+c*2.35,3.18+r*.92,1.8,.28,12,MUTED)
    oval(s,8.05,1.48,5.25,5.48,SOFT_TEAL); native_monitor(s,8.62,2.20,1.55,.95)
    add_rect(s,10.65,2.22,.42,1.55,WHITE,TEAL,True,1); add_circle(s,10.73,2.42,.25,MIST,NAVY,1)
    add_line(s,10.86,3.78,10.36,4.85,NAVY,1.4); add_line(s,10.86,3.78,11.52,4.85,NAVY,1.4)
    add_circle(s,9.05,5.22,.55,WHITE,ORANGE,1.3); add_circle(s,11.32,5.22,.55,WHITE,TEAL,1.3)
    add_text(s,"視訊・裂隙燈・皮膚鏡概念",8.60,6.02,3.70,.30,10,MUTED,False,PP_ALIGN.CENTER)
    notes(s,"IL-04 以原生線條呈現視訊與設備概念，不宣稱實際已提供科別。")

    # 10 outcomes story
    s=page(prs,10,"成果應以一致口徑的 KPI 集中檢視","Outcomes")
    oval(s,-1.10,1.35,6.45,6.35,SOFT_ORANGE)
    pill(s,"核心成果",.78,1.72,1.10,WHITE,ORANGE); add_text(s,"〔待補資料〕",.78,2.60,4.10,.72,37,NAVY,True); add_text(s,"服務量",.80,3.52,2.0,.32,14,TEXT,True); add_text(s,"定義與統計期間〔待補資料〕",.80,4.08,3.5,.40,11,MUTED)
    add_text(s,"一個主結果，\n五個支持證據。",5.72,1.66,3.1,.82,22,NAVY,True)
    for i,(lab,c) in enumerate([("合作院所",TEAL),("服務科別",NAVY),("完成率",ORANGE),("回應時效",TEAL),("滿意度",NAVY)]):
        x=5.75+(i%2)*3.25; y=3.02+(i//2)*1.20; metric(s,x,y,lab,16,c,"期間／定義〔待補〕")
    notes(s,"Hero KPI storytelling；保留全部待補資料，不使用六張同款卡。")

    # 11 full-width chart composition
    s=page(prs,11,"服務使用趨勢需搭配事件與統計口徑解讀","Trend")
    add_text(s,"趨勢不是一條線，\n而是服務演進的證據。",.72,1.50,4.35,.82,21,NAVY,True)
    pill(s,"原始時序資料〔待補〕",9.65,1.58,2.55,SOFT_ORANGE,ORANGE)
    add_rect(s,.72,2.55,11.88,3.82,WHITE,None)
    for i in range(4): add_line(s,1.05,3.15+i*.72,12.10,3.15+i*.72,"E3EDF2",1)
    add_line(s,1.05,5.73,12.10,5.73,LINE,1.2)
    # A polished placeholder path—explicitly not data.
    pts=[(1.25,5.35),(3.10,4.92),(4.82,5.08),(6.62,4.20),(8.55,4.42),(10.25,3.52),(11.82,3.78)]
    for i in range(len(pts)-1): add_line(s,*pts[i],*pts[i+1],TEAL,2)
    for x,y in pts: add_circle(s,x-.09,y-.09,.18,WHITE,TEAL,1)
    add_text(s,"視覺結構 placeholder｜不代表實際趨勢",1.08,5.92,4.1,.25,9,MUTED)
    add_line(s,8.52,2.78,8.52,5.72,ORANGE,1,True); pill(s,"關鍵事件〔待補〕",8.72,2.82,1.65,SOFT_ORANGE,ORANGE)
    notes(s,"正式 full-width chart composition；線形僅為清楚標記的視覺 placeholder。")

    # 12 heatmap
    s=page(prs,12,"服務分布可從地區、院所、科別或時段辨識差異","Distribution")
    add_text(s,"服務分布熱力表",.72,1.52,3.0,.36,16,NAVY,True); add_text(s,"維度、分母與原始值〔待補資料〕",.72,1.95,4.0,.28,10,MUTED)
    colors=["F3F8FA","E4F0F1","CFE5E4","A8D1CF","78B6B2"]
    for r in range(5):
        add_text(s,f"維度 {r+1}",.78,2.63+r*.66,1.15,.24,10,MUTED)
        for c in range(8): add_rect(s,2.05+c*.76,2.48+r*.66,.62,.46,colors[(r*2+c)%5],WHITE,False)
    add_text(s,"低",2.05,6.00,.4,.20,9,MUTED); add_text(s,"高",7.45,6.00,.4,.20,9,MUTED,False,PP_ALIGN.RIGHT)
    add_text(s,"洞察",9.12,1.60,1.0,.28,11,ORANGE,True); add_text(s,"分布差異\n〔待補資料〕",9.12,2.25,3.05,.78,22,NAVY,True); add_line(s,9.12,3.35,11.95,3.35,PALE,1); add_text(s,"解讀必須同時呈現分母、統計期間與服務情境。",9.12,3.78,3.0,1.05,13,TEXT)
    notes(s,"正式 heatmap layout；色階為結構 placeholder，不代表服務量。")

    # 13 editorial case + IL05
    s=page(prs,13,"特色亮點必須同時呈現做法與可核實價值","Highlights")
    oval(s,7.45,1.25,6.20,6.10,SOFT_TEAL); native_person(s,8.18,2.45,1.05,True); native_person(s,10.82,3.15,.92,False)
    for yy in [3.10,3.34,3.58]: add_line(s,9.38,yy,10.70,yy,ORANGE if yy==3.34 else TEAL,1.2,True)
    add_text(s,"專科醫療，\n延伸到需要的地方。",.72,1.65,5.65,1.20,27,NAVY,True)
    add_text(s,"亮點做法 01",.74,3.35,1.7,.28,11,TEAL,True); add_text(s,"做法與佐證〔待補資料〕",.74,3.82,4.8,.38,15,TEXT,True)
    add_line(s,.74,4.48,5.55,4.48,PALE,1)
    add_text(s,"公開案例",.74,4.82,1.7,.28,11,ORANGE,True); add_text(s,"情境與成效〔待補資料〕",.74,5.28,4.8,.38,15,TEXT,True)
    notes(s,"IL-05 以兩名原生線性人物與照護連線呈現人本連結。")

    # 14 problem to action
    s=page(prs,14,"已知挑戰需要對應具體改善行動與責任","Challenges")
    add_text(s,"從問題走向可追蹤的行動",.72,1.55,5.2,.42,17,NAVY,True)
    for i in range(3):
        y=2.38+i*1.30
        add_text(s,f"0{i+1}",.80,y,.45,.30,11,ORANGE,True); add_text(s,"挑戰〔待補資料〕",1.42,y,2.60,.34,14,NAVY,True)
        add_line(s,4.12,y+.16,5.12,y+.16,TEAL,1.4); add_text(s,"改善行動〔待補資料〕",5.38,y,3.08,.34,14,TEXT,True); pill(s,"Owner／期限〔待補〕",9.40,y-.02,2.58,SOFT_BLUE,NAVY)
        add_line(s,.82,y+.72,12.00,y+.72,PALE,1)
    oval(s,9.92,5.72,4.0,1.75,SOFT_ORANGE); add_text(s,"排序原則｜影響 × 急迫 × 可行",8.82,6.22,3.65,.28,10,ORANGE,True,PP_ALIGN.CENTER)
    notes(s,"Problem → Action composition，降低框線密度。")

    # 15 roadmap + IL06
    s=page(prs,15,"下一階段以可執行路線圖擴大服務與影響","Next stage")
    add_text(s,"從穩定服務，走向網絡深化。",.72,1.55,5.2,.42,18,NAVY,True)
    add_line(s,.92,3.78,7.55,3.78,TEAL,2)
    for i,(stage,title,c) in enumerate([("近期","穩定與補強",NAVY),("中期","擴展與整合",TEAL),("後期","深化與評估",ORANGE)]):
        x=1.05+i*2.65; add_circle(s,x,3.40,.76,MIST,c,2); pill(s,stage,x-.04,2.34,.84,WHITE,c); add_text(s,title,x-.30,4.55,1.45,.32,13,NAVY,True,PP_ALIGN.CENTER); add_text(s,"行動／時程\n〔待補資料〕",x-.48,5.05,1.85,.60,10,MUTED,False,PP_ALIGN.CENTER)
    oval(s,8.05,1.32,5.55,5.82,SOFT_TEAL); native_network(s,10.72,3.72,1.65,7); add_text(s,"成功指標〔待補資料〕",9.12,6.22,3.25,.30,11,ORANGE,True,PP_ALIGN.CENTER)
    notes(s,"IL-06 以原生中心醫院與多服務據點網絡收束全套簡報。")

    prs.save(PPTX); print(PPTX)


if __name__ == "__main__": build()

